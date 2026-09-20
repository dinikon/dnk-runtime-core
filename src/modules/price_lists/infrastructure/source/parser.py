from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterator

import yaml
from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook

from src.modules.price_lists.domain import (
    MappingValidationError,
    canonical_state_hash,
    normalize_availability,
)

PROM_XML_SOURCE = {"item_path": "yml_catalog.shop.offers.offer"}
PROM_XML_MAPPING = {
    "external_id": {"selector": "@id", "required": True, "trim": True},
    "sku": {"selector": "vendorCode", "required": True, "trim": True},
    "title": {"selector": "name", "required": True, "trim": True},
    "purchase_price": {"selector": "price", "type": "decimal", "required": True},
    "rrp": {"selector": "priceRRP", "type": "decimal"},
    "currency": {"selector": "currencyId", "default": "UAH"},
    "availability": {
        "selector": "@available",
        "default": "out_of_stock",
        "map": {
            "склад": "in_stock",
            "true": "in_stock",
            "false": "out_of_stock",
            "": "out_of_stock",
        },
    },
    "quantity": {"constant": None},
}


def prom_xml_config() -> tuple[dict[str, Any], dict[str, Any]]:
    return dict(PROM_XML_SOURCE), {
        key: dict(value) for key, value in PROM_XML_MAPPING.items()
    }


@dataclass(frozen=True, slots=True)
class ParsedRow:
    row_number: int
    normalized: dict[str, Any]
    errors: tuple[str, ...]


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _path_value(value: Any, selector: str) -> Any:
    current = value
    for part in selector.split("."):
        if isinstance(current, dict):
            current = current.get(part)
        else:
            return None
    return current


def _extract(raw: Any, specification: dict[str, Any]) -> Any:
    if "constant" in specification:
        return specification["constant"]
    selectors = specification.get("selectors") or [specification.get("selector")]
    for selector in selectors:
        if not selector:
            continue
        if isinstance(raw, dict):
            if isinstance(selector, int):
                row_values = list(raw.values())
                value = row_values[selector] if selector < len(row_values) else None
            else:
                value = _path_value(raw, str(selector))
        else:
            if str(selector).startswith("@"):
                value = raw.attrib.get(str(selector)[1:])
            else:
                node = raw
                for part in str(selector).split("."):
                    node = next(
                        (child for child in node if _local_name(child.tag) == part),
                        None,
                    )
                    if node is None:
                        break
                value = node.text if node is not None else None
        if value not in (None, ""):
            return value
    return specification.get("default")


def _decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    normalized = re.sub(r"\s+", "", str(value)).replace(",", ".")
    try:
        result = Decimal(normalized)
    except InvalidOperation as exc:
        raise ValueError("must be a decimal") from exc
    if result < 0:
        raise ValueError("must not be negative")
    return result


def _integer(value: Any) -> int | None:
    if value in (None, ""):
        return None
    result = int(Decimal(str(value).replace(",", ".")))
    if result < 0:
        raise ValueError("must not be negative")
    return result


def normalize_row(
    raw: Any, mapping: dict[str, Any]
) -> tuple[dict[str, Any], tuple[str, ...]]:
    values: dict[str, Any] = {}
    errors: list[str] = []
    for field, specification in mapping.items():
        specification = specification or {}
        try:
            value = _extract(raw, specification)
            if isinstance(value, str) and specification.get("trim", True):
                value = value.strip()
            replacements = specification.get("replace") or {}
            if isinstance(value, str) and isinstance(replacements, dict):
                for old, new in replacements.items():
                    value = value.replace(str(old), str(new))
            if isinstance(value, str) and specification.get("lower"):
                value = value.lower()
            lookup = specification.get("map") or {}
            if isinstance(lookup, dict):
                value = lookup.get(value, lookup.get(str(value).casefold(), value))
            if specification.get("type") == "decimal" or field in {
                "purchase_price",
                "rrp",
            }:
                value = _decimal(value)
            elif specification.get("type") == "integer" or field == "quantity":
                value = _integer(value)
            if specification.get("required") and value in (None, ""):
                errors.append(f"{field}: required")
            values[field] = value
        except (ValueError, TypeError) as exc:
            errors.append(f"{field}: {exc}")
            values[field] = None
    for required in ("external_id", "sku", "title", "purchase_price"):
        if values.get(required) in (None, "") and not any(
            error.startswith(f"{required}:") for error in errors
        ):
            errors.append(f"{required}: required")
    currency = str(values.get("currency") or "").strip().upper()
    if not re.fullmatch(r"[A-Z]{3}", currency):
        errors.append("currency: expected ISO 4217 code")
    values["currency"] = currency
    quantity = values.get("quantity")
    values["availability"] = normalize_availability(
        values.get("availability"), quantity
    )
    if not errors:
        values["value_hash"] = canonical_state_hash(
            purchase_price=values["purchase_price"],
            rrp=values.get("rrp"),
            currency=currency,
            availability=values["availability"],
            quantity=quantity,
        )
    return values, tuple(errors)


class SourceParser:
    max_rows = 500_000
    max_columns = 256
    max_uncompressed_bytes = 256 * 1024 * 1024

    def inspect(
        self, path: Path, source_format: str, source_config: dict[str, Any]
    ) -> dict[str, Any]:
        if source_format == "xlsx":
            self._validate_xlsx_archive(path)
            # Remote files are intentionally stored under random extensionless
            # temporary names. Passing a binary stream makes openpyxl validate
            # the ZIP payload instead of rejecting the temporary filename.
            with path.open("rb") as source:
                workbook = load_workbook(source, read_only=True, data_only=True)
                try:
                    sheets = list(workbook.sheetnames)
                    sheet = workbook[source_config.get("sheet_name") or sheets[0]]
                    header_row = int(source_config.get("header_row", 1))
                    rows = sheet.iter_rows(
                        min_row=header_row, max_row=header_row, values_only=True
                    )
                    headers = [
                        str(value).strip() for value in next(rows) if value is not None
                    ]
                    return {"sheets": sheets, "columns": headers}
                finally:
                    workbook.close()
        if source_format == "xml":
            paths: set[str] = set()
            stack: list[str] = []
            for event, node in iterparse(path, events=("start", "end")):
                if event == "start":
                    stack.append(_local_name(node.tag))
                    if len(stack) > 64:
                        raise MappingValidationError("XML nesting limit exceeded.")
                    if len(paths) < 100:
                        paths.add(".".join(stack))
                else:
                    stack.pop()
                    node.clear()
            return {"paths": sorted(paths)}
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
        return {"paths": self._yaml_paths(data)}

    def rows(
        self,
        path: Path,
        source_format: str,
        source_config: dict[str, Any],
        mapping: dict[str, Any],
        *,
        limit: int | None = None,
    ) -> Iterator[ParsedRow]:
        raw_rows: Iterator[tuple[int, Any]]
        if source_format == "xml":
            raw_rows = self._xml_rows(path, source_config)
        elif source_format == "yaml":
            raw_rows = self._yaml_rows(path, source_config)
        elif source_format == "xlsx":
            raw_rows = self._xlsx_rows(path, source_config)
        else:
            raise MappingValidationError("Unsupported source format.")
        for index, (row_number, raw) in enumerate(raw_rows):
            if limit is not None and index >= limit:
                break
            values, errors = normalize_row(raw, mapping)
            yield ParsedRow(row_number=row_number, normalized=values, errors=errors)

    def _xml_rows(
        self, path: Path, source_config: dict[str, Any]
    ) -> Iterator[tuple[int, Any]]:
        item_path = str(source_config.get("item_path") or "")
        target_path = [part for part in item_path.split(".") if part]
        if not target_path:
            raise MappingValidationError("XML item_path is required.")
        count = 0
        stack: list[str] = []
        for event, node in iterparse(path, events=("start", "end")):
            if event == "start":
                stack.append(_local_name(node.tag))
                if len(stack) > 64:
                    raise MappingValidationError("XML nesting limit exceeded.")
                continue
            if stack == target_path:
                count += 1
                if count > self.max_rows:
                    raise MappingValidationError("Source row limit exceeded.")
                yield count, node
                node.clear()
            stack.pop()

    def _yaml_rows(
        self, path: Path, source_config: dict[str, Any]
    ) -> Iterator[tuple[int, Any]]:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
        item_path = str(source_config.get("item_path") or "")
        selected = _path_value(data, item_path) if item_path else data
        if not isinstance(selected, list):
            raise MappingValidationError("YAML item_path must select a list.")
        for index, row in enumerate(selected, start=1):
            if index > self.max_rows:
                raise MappingValidationError("Source row limit exceeded.")
            if not isinstance(row, dict):
                yield index, {}
            else:
                yield index, row

    def _xlsx_rows(
        self, path: Path, source_config: dict[str, Any]
    ) -> Iterator[tuple[int, Any]]:
        self._validate_xlsx_archive(path)
        with path.open("rb") as source:
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                sheet_name = source_config.get("sheet_name") or workbook.sheetnames[0]
                if sheet_name not in workbook.sheetnames:
                    raise MappingValidationError(
                        "Configured XLSX sheet does not exist."
                    )
                sheet = workbook[sheet_name]
                header_row = int(source_config.get("header_row", 1))
                data_start = int(source_config.get("data_start_row", header_row + 1))
                header_values = next(
                    sheet.iter_rows(
                        min_row=header_row, max_row=header_row, values_only=True
                    )
                )
                headers = [
                    str(value).strip() if value is not None else ""
                    for value in header_values
                ]
                if len(headers) > self.max_columns:
                    raise MappingValidationError("XLSX column limit exceeded.")
                for count, values in enumerate(
                    sheet.iter_rows(min_row=data_start, values_only=True), start=0
                ):
                    if count >= self.max_rows:
                        raise MappingValidationError("Source row limit exceeded.")
                    if all(value is None for value in values):
                        continue
                    yield data_start + count, dict(zip(headers, values, strict=False))
            finally:
                workbook.close()

    def _validate_xlsx_archive(self, path: Path) -> None:
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 10_000:
                    raise MappingValidationError("XLSX archive entry limit exceeded.")
                if (
                    sum(entry.file_size for entry in entries)
                    > self.max_uncompressed_bytes
                ):
                    raise MappingValidationError(
                        "XLSX uncompressed size limit exceeded."
                    )
        except zipfile.BadZipFile as exc:
            raise MappingValidationError("Invalid XLSX archive.") from exc

    def _yaml_paths(self, value: Any, prefix: str = "") -> list[str]:
        found: list[str] = []
        if isinstance(value, dict):
            for key, child in list(value.items())[:100]:
                path = f"{prefix}.{key}" if prefix else str(key)
                found.append(path)
                found.extend(self._yaml_paths(child, path))
        return found[:100]
